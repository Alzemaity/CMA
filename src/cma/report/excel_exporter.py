import datetime
import logging
import pathlib

logger = logging.getLogger(__name__)

REPORTS_DIR = pathlib.Path("reports")

# Column definitions: (Arabic header, dict key)
COLUMNS = [
    ("رقم", None),               # auto index
    ("نوع الوحدة", "نوع_الوحدة"),
    ("نوع العرض", "نوع_العرض"),
    ("السعر (ريال)", "السعر"),
    ("المساحة (م²)", "المساحة_متر"),
    ("سعر المتر (ريال)", "السعر_للمتر"),
    ("المدينة", "المدينة"),
    ("الحي", "الحي"),
    ("الشارع", "الشارع"),
    ("المنطقة", "المنطقة"),
    ("الفئة", "الفئة"),
    ("غرف النوم", "غرف_النوم"),
    ("الصالات", "الصالات"),
    ("دورات المياه", "دورات_المياه"),
    ("الدور", "الدور"),
    ("عمر العقار (سنة)", "عمر_العقار"),
    ("المميزات", "المميزات"),
    ("رقم التواصل", "رقم_التواصل"),
    ("تاريخ الإعلان", "تاريخ_الإعلان"),
    ("رابط الإعلان", "رابط_الإعلان"),
]


class ExcelExporter:
    def export(
        self,
        listings: list[dict],
        city: str,
        neighborhood: str,
        report_id: int,
    ) -> str | None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (
                Alignment,
                Border,
                Font,
                PatternFill,
                Side,
            )
            from openpyxl.utils import get_column_letter

            REPORTS_DIR.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"cma_data_{city}_{neighborhood}_{timestamp}.xlsx"
            excel_path = REPORTS_DIR / filename

            wb = Workbook()
            ws = wb.active
            ws.title = "بيانات الإعلانات"
            ws.sheet_view.rightToLeft = True  # RTL sheet

            # ── Styles ────────────────────────────────────────────────────────
            header_fill = PatternFill(
                start_color="1A1A2E", end_color="1A1A2E", fill_type="solid"
            )
            header_font = Font(
                name="Calibri", bold=True, color="FFFFFF", size=11
            )
            alt_fill = PatternFill(
                start_color="F0F4FF", end_color="F0F4FF", fill_type="solid"
            )
            link_font = Font(
                name="Calibri", color="0563C1", underline="single", size=10
            )
            center = Alignment(
                horizontal="center", vertical="center",
                wrap_text=True,
            )
            right_align = Alignment(
                horizontal="right", vertical="center",
                wrap_text=True,
            )
            thin = Side(style="thin", color="D0D7DE")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            # ── Title row ────────────────────────────────────────────────────
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
            title_cell = ws.cell(row=1, column=1)
            title_cell.value = (
                f"بيانات إعلانات {city} - حي {neighborhood}  |  "
                f"التقرير #{report_id}  |  {datetime.date.today().isoformat()}"
            )
            title_cell.font = Font(name="Calibri", bold=True, size=14, color="1A1A2E")
            title_cell.alignment = center
            title_cell.fill = PatternFill(
                start_color="E8EAF6", end_color="E8EAF6", fill_type="solid"
            )
            ws.row_dimensions[1].height = 28

            # ── Sub-info row ─────────────────────────────────────────────────
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
            sub_cell = ws.cell(row=2, column=1)
            sub_cell.value = (
                f"إجمالي الإعلانات: {len(listings)}  |  "
                f"المرحلة الأولى: بيانات خام بدون تقييم  |  "
                f"يمكنك استخدام فلاتر إكسيل للبحث والتحليل"
            )
            sub_cell.font = Font(name="Calibri", size=10, color="555555", italic=True)
            sub_cell.alignment = center
            ws.row_dimensions[2].height = 18

            # ── Header row ────────────────────────────────────────────────────
            HEADER_ROW = 3
            for col_idx, (col_header, _) in enumerate(COLUMNS, start=1):
                cell = ws.cell(row=HEADER_ROW, column=col_idx, value=col_header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = border
            ws.row_dimensions[HEADER_ROW].height = 22

            # ── Data rows ─────────────────────────────────────────────────────
            for row_idx, listing in enumerate(listings, start=1):
                excel_row = HEADER_ROW + row_idx
                is_alt = row_idx % 2 == 0

                for col_idx, (_, key) in enumerate(COLUMNS, start=1):
                    cell = ws.cell(row=excel_row, column=col_idx)
                    cell.border = border
                    cell.alignment = right_align

                    if is_alt:
                        cell.fill = alt_fill

                    if key is None:
                        # Auto-index column
                        cell.value = row_idx
                        cell.alignment = center
                    elif key == "رابط_الإعلان":
                        url = listing.get(key, "")
                        if url:
                            cell.value = url
                            cell.hyperlink = url
                            cell.font = link_font
                        else:
                            cell.value = "-"
                    elif key == "المميزات":
                        amenities = listing.get(key, [])
                        if isinstance(amenities, list):
                            cell.value = " | ".join(amenities) if amenities else "-"
                        else:
                            cell.value = str(amenities) if amenities else "-"
                    else:
                        val = listing.get(key)
                        if val is None:
                            cell.value = "-"
                        elif isinstance(val, float):
                            cell.value = round(val, 2)
                        else:
                            cell.value = val

                ws.row_dimensions[excel_row].height = 20

            # ── Column widths ─────────────────────────────────────────────────
            col_widths = [
                5,   # رقم
                12,  # نوع الوحدة
                10,  # نوع العرض
                14,  # السعر
                12,  # المساحة
                14,  # سعر المتر
                12,  # المدينة
                14,  # الحي
                16,  # الشارع
                14,  # المنطقة
                10,  # الفئة
                10,  # غرف النوم
                8,   # الصالات
                12,  # دورات المياه
                8,   # الدور
                14,  # عمر العقار
                30,  # المميزات
                16,  # رقم التواصل
                14,  # تاريخ الإعلان
                45,  # رابط الإعلان
            ]
            for i, width in enumerate(col_widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = width

            # ── Freeze header ─────────────────────────────────────────────────
            ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

            # ── Auto-filter ───────────────────────────────────────────────────
            ws.auto_filter.ref = (
                f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}{HEADER_ROW + len(listings)}"
            )

            wb.save(str(excel_path))
            logger.info(f"Excel exported: {excel_path} ({len(listings)} rows)")
            return str(excel_path)

        except ImportError:
            logger.warning("openpyxl not available, skipping Excel export")
            return None
        except Exception as e:
            logger.error(f"Excel export failed: {e}")
            return None
